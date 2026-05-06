from django.shortcuts import render
import base64 as b64lib
import random
import string
import json
import urllib.parse


def texttools_home(request):
    return render(request, 'texttools/home.html')


def word_counter(request):
    result = None
    text = ''
    if request.method == 'POST':
        text = request.POST.get('text', '')
        words = len(text.split()) if text.strip() else 0
        chars = len(text)
        chars_no_space = len(text.replace(' ', ''))
        sentences = text.count('.') + text.count('!') + text.count('?')
        paragraphs = len([p for p in text.split('\n') if p.strip()])
        result = {
            'words': words,
            'chars': chars,
            'chars_no_space': chars_no_space,
            'sentences': sentences,
            'paragraphs': paragraphs,
            'reading_time': max(1, round(words / 200)),
        }
    return render(request, 'texttools/word_counter.html', {'result': result, 'text': text})


def password_generator(request):
    result = None
    length = 16
    if request.method == 'POST':
        length = min(max(int(request.POST.get('length', 16)), 6), 64)
        use_upper = 'uppercase' in request.POST
        use_lower = 'lowercase' in request.POST
        use_numbers = 'numbers' in request.POST
        use_symbols = 'symbols' in request.POST
        pool = ''
        if use_upper:   pool += string.ascii_uppercase
        if use_lower:   pool += string.ascii_lowercase
        if use_numbers: pool += string.digits
        if use_symbols: pool += '!@#$%^&*()_+-=[]{}|;:,.<>?'
        if not pool:    pool = string.ascii_letters + string.digits
        result = ''.join(random.choices(pool, k=length))
    return render(request, 'texttools/password_generator.html', {'result': result, 'length': length})


CASE_OPTIONS = [
    ('upper', 'UPPERCASE'),
    ('lower', 'lowercase'),
    ('title', 'Title Case'),
    ('sentence', 'Sentence case'),
    ('alternate', 'aLtErNaTe'),
    ('reverse', 'esreveR'),
]


def text_case_converter(request):
    result = None
    text = ''
    case = 'upper'
    if request.method == 'POST':
        text = request.POST.get('text', '')
        case = request.POST.get('case', 'upper')
        if case == 'upper':
            result = text.upper()
        elif case == 'lower':
            result = text.lower()
        elif case == 'title':
            result = text.title()
        elif case == 'sentence':
            result = '. '.join(s.strip().capitalize() for s in text.split('.'))
        elif case == 'alternate':
            result = ''.join(c.upper() if i % 2 == 0 else c.lower() for i, c in enumerate(text))
        elif case == 'reverse':
            result = text[::-1]
    return render(request, 'texttools/text_case.html', {
        'result': result, 'text': text, 'case': case, 'cases': CASE_OPTIONS
    })


_LOREM_SENTENCES = [
    "Lorem ipsum dolor sit amet, consectetur adipiscing elit.",
    "Sed do eiusmod tempor incididunt ut labore et dolore magna aliqua.",
    "Ut enim ad minim veniam, quis nostrud exercitation ullamco laboris.",
    "Duis aute irure dolor in reprehenderit in voluptate velit esse cillum.",
    "Excepteur sint occaecat cupidatat non proident, sunt in culpa qui officia.",
    "Curabitur pretium tincidunt lacus, nulla facilisis nisl facilisis.",
    "Nunc molestie augue vel facilisis varius, lorem massa fermentum.",
    "Pellentesque habitant morbi tristique senectus et netus et malesuada.",
    "Vestibulum ante ipsum primis in faucibus orci luctus et ultrices.",
    "Fusce dapibus tellus ac cursus commodo, tortor mauris condimentum nibh.",
    "Integer nec odio praesent libero sed cursus ante dapibus diam.",
    "Nulla quis sem at nibh elementum imperdiet duis sagittis ipsum.",
]

_LOREM_PARA = (
    "Lorem ipsum dolor sit amet, consectetur adipiscing elit. Sed do eiusmod tempor incididunt "
    "ut labore et dolore magna aliqua. Ut enim ad minim veniam, quis nostrud exercitation ullamco "
    "laboris nisi ut aliquip ex ea commodo consequat. Duis aute irure dolor in reprehenderit in "
    "voluptate velit esse cillum dolore eu fugiat nulla pariatur. Excepteur sint occaecat cupidatat "
    "non proident, sunt in culpa qui officia deserunt mollit anim id est laborum."
)


def lorem_ipsum(request):
    result = None
    if request.method == 'POST':
        count = min(max(int(request.POST.get('count', 3)), 1), 20)
        unit = request.POST.get('type', 'paragraphs')
        if unit == 'paragraphs':
            result = [_LOREM_PARA for _ in range(count)]
        elif unit == 'words':
            words = (_LOREM_PARA.split() * 10)[:count]
            result = [' '.join(words)]
        else:
            sentences = (_LOREM_SENTENCES * 5)[:count]
            result = sentences
    return render(request, 'texttools/lorem_ipsum.html', {'result': result})


def base64_tool(request):
    result = None
    text = ''
    action = 'encode'
    error = None
    if request.method == 'POST':
        text = request.POST.get('text', '')
        action = request.POST.get('action', 'encode')
        if not text.strip():
            error = 'Please enter some text to convert.'
        else:
            try:
                if action == 'encode':
                    result = b64lib.b64encode(text.encode('utf-8')).decode('utf-8')
                else:
                    result = b64lib.b64decode(text.encode('utf-8')).decode('utf-8')
            except Exception:
                error = 'Invalid Base64 input. Please enter valid Base64 encoded text.'
    return render(request, 'texttools/base64.html', {'result': result, 'text': text, 'action': action, 'error': error})


def json_formatter(request):
    result = None
    text = ''
    error = None
    if request.method == 'POST':
        text = request.POST.get('text', '')
        try:
            parsed = json.loads(text)
            result = json.dumps(parsed, indent=2, ensure_ascii=False)
        except json.JSONDecodeError as e:
            error = f'Invalid JSON: {e}'
    return render(request, 'texttools/json_formatter.html', {'result': result, 'text': text, 'error': error})


def url_encoder(request):
    result = None
    text = ''
    action = 'encode'
    if request.method == 'POST':
        text = request.POST.get('text', '')
        action = request.POST.get('action', 'encode')
        if action == 'encode':
            result = urllib.parse.quote(text, safe='')
        else:
            result = urllib.parse.unquote(text)
    return render(request, 'texttools/url_encoder.html', {'result': result, 'text': text, 'action': action})


def remove_duplicates(request):
    result = None
    text = ''
    count_removed = 0
    if request.method == 'POST':
        text = request.POST.get('text', '')
        lines = text.splitlines()
        seen = set()
        out = []
        for line in lines:
            line = line.rstrip()
            if line not in seen:
                seen.add(line)
                out.append(line)
        count_removed = len(lines) - len(out)
        result = '\n'.join(out)
    return render(request, 'texttools/remove_duplicates.html', {
        'result': result, 'text': text, 'count_removed': count_removed
    })


# ──────────── QR Code Generator ────────────
def qr_generator(request):
    if request.method == 'POST':
        try:
            import qrcode
            import io
            from django.http import FileResponse
            mode = request.POST.get('mode', 'single')
            if mode == 'bulk' and request.FILES.get('file'):
                import openpyxl, zipfile, tempfile, os, uuid, re
                wb = openpyxl.load_workbook(request.FILES['file'], data_only=True)
                ws = wb.active
                buf = io.BytesIO()
                with zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED) as zf:
                    n = 0
                    for r_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
                        for c_idx, val in enumerate(row, start=1):
                            if val is None or str(val).strip() == '':
                                continue
                            img = qrcode.make(str(val))
                            img_buf = io.BytesIO()
                            img.save(img_buf, format='PNG')
                            safe = re.sub(r'[^A-Za-z0-9._-]+', '_', str(val))[:40] or 'qr'
                            zf.writestr(f'qr_r{r_idx}_c{c_idx}_{safe}.png', img_buf.getvalue())
                            n += 1
                buf.seek(0)
                resp = FileResponse(buf, as_attachment=True, filename='qr_codes.zip', content_type='application/zip')
                return resp
            content = request.POST.get('content', '').strip()
            if not content:
                return render(request, 'texttools/qr_generator.html', {'error': 'Please enter some content.'})
            box_size = int(request.POST.get('box_size', 10))
            border = int(request.POST.get('border', 4))
            ec_map = {'L': qrcode.constants.ERROR_CORRECT_L, 'M': qrcode.constants.ERROR_CORRECT_M,
                      'Q': qrcode.constants.ERROR_CORRECT_Q, 'H': qrcode.constants.ERROR_CORRECT_H}
            ec = ec_map.get(request.POST.get('ec', 'M'), qrcode.constants.ERROR_CORRECT_M)
            qr = qrcode.QRCode(version=None, error_correction=ec, box_size=box_size, border=border)
            qr.add_data(content)
            qr.make(fit=True)
            img = qr.make_image(fill_color='black', back_color='white')
            buf = io.BytesIO()
            img.save(buf, format='PNG')
            buf.seek(0)
            data_url = 'data:image/png;base64,' + b64lib.b64encode(buf.getvalue()).decode()
            return render(request, 'texttools/qr_generator.html', {'image': data_url, 'content': content})
        except Exception as e:
            return render(request, 'texttools/qr_generator.html', {'error': str(e)})
    return render(request, 'texttools/qr_generator.html')


# ──────────── Barcode Generator ────────────
def barcode_generator(request):
    if request.method == 'POST':
        try:
            import barcode
            from barcode.writer import ImageWriter
            import io
            from django.http import FileResponse
            mode = request.POST.get('mode', 'single')
            fmt = request.POST.get('format', 'code128').lower()

            def make_one(value, fmt):
                value = str(value)
                try:
                    cls = barcode.get_barcode_class(fmt)
                    bc = cls(value, writer=ImageWriter())
                except Exception:
                    cls = barcode.get_barcode_class('code128')
                    bc = cls(value, writer=ImageWriter())
                buf = io.BytesIO()
                bc.write(buf)
                buf.seek(0)
                return buf.getvalue()

            if mode == 'bulk' and request.FILES.get('file'):
                import openpyxl, zipfile, re
                wb = openpyxl.load_workbook(request.FILES['file'], data_only=True)
                ws = wb.active
                out = io.BytesIO()
                with zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED) as zf:
                    for r_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
                        for c_idx, val in enumerate(row, start=1):
                            if val is None or str(val).strip() == '':
                                continue
                            png = make_one(val, fmt)
                            safe = re.sub(r'[^A-Za-z0-9._-]+', '_', str(val))[:40] or 'bc'
                            zf.writestr(f'barcode_r{r_idx}_c{c_idx}_{safe}.png', png)
                out.seek(0)
                return FileResponse(out, as_attachment=True, filename='barcodes.zip', content_type='application/zip')

            content = request.POST.get('content', '').strip()
            if not content:
                return render(request, 'texttools/barcode_generator.html', {'error': 'Please enter content.'})
            png = make_one(content, fmt)
            data_url = 'data:image/png;base64,' + b64lib.b64encode(png).decode()
            return render(request, 'texttools/barcode_generator.html', {'image': data_url, 'content': content, 'format': fmt})
        except Exception as e:
            return render(request, 'texttools/barcode_generator.html', {'error': str(e)})
    return render(request, 'texttools/barcode_generator.html')
